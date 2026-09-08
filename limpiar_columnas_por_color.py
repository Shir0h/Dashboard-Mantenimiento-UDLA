"""
limpiar_columnas_por_color.py

Lee un archivo Excel y conserva solo las columnas cuyo encabezado esté marcado
con color verde o anaranjado. Genera salida en outputs/ y log en logs/.

Uso:
    python limpiar_columnas_por_color.py

Uso indicando archivo y hoja:
    python limpiar_columnas_por_color.py --input data/raw/Data_proyecto_integrador.xlsx --sheet Hoja1
"""

from pathlib import Path
from datetime import datetime
import argparse
import pandas as pd
from openpyxl import load_workbook
from loguru import logger

try:
    from src.config_utils import cargar_config, resolver_ruta
    from src.logging_utils import configurar_logging
except ModuleNotFoundError:
    from config_utils import cargar_config, resolver_ruta  # type: ignore
    from logging_utils import configurar_logging  # type: ignore


CONFIG = cargar_config()
INPUT_DEFAULT = resolver_ruta(CONFIG.get("paths", {}).get("input_excel"), "data/raw/Data_proyecto_integrador.xlsx")
OUTPUT_DIR = resolver_ruta(CONFIG.get("paths", {}).get("output_dir"), "outputs")
LOG_DIR = resolver_ruta(CONFIG.get("paths", {}).get("log_dir"), "logs")

GREEN_HINTS = set(CONFIG.get("processing", {}).get("green_hints", []))
ORANGE_HINTS = set(CONFIG.get("processing", {}).get("orange_hints", []))


def configurar_logs() -> Path:
    """Inicializa los logs con un formato más claro y visual para consola y archivo."""
    return configurar_logging(LOG_DIR, "limpieza_columnas_color")


def normalizar_rgb(rgb: str | None) -> str:
    """Convierte un color a un formato hexadecimal limpio para compararlo con facilidad."""
    if rgb is None:
        return ""
    rgb = str(rgb).upper().replace("#", "")
    if len(rgb) == 8:
        rgb = rgb[-6:]
    return rgb


def es_color_objetivo(rgb: str) -> bool:
    """Revisa si un color parece verde o naranja según reglas simples y códigos conocidos."""
    rgb = normalizar_rgb(rgb)

    if not rgb:
        return False

    if rgb in GREEN_HINTS or rgb in ORANGE_HINTS:
        return True

    try:
        r = int(rgb[0:2], 16)
        g = int(rgb[2:4], 16)
        b = int(rgb[4:6], 16)
    except ValueError:
        return False

    es_verde = g >= 140 and g > r and g > b
    es_anaranjado = r >= 200 and 80 <= g <= 210 and b <= 120

    return es_verde or es_anaranjado


def obtener_color_celda(cell) -> str:
    """Lee el color de fondo de una celda para saber si su encabezado debe conservarse."""
    fill = cell.fill

    if fill is None or fill.fill_type is None:
        return ""

    fg_color = fill.fgColor

    if fg_color.type == "rgb":
        return normalizar_rgb(fg_color.rgb)

    if fg_color.type == "indexed":
        return f"INDEXED_{fg_color.indexed}"

    if fg_color.type == "theme":
        return f"THEME_{fg_color.theme}"

    return ""


def detectar_columnas_por_color(excel_path: Path, sheet_name: str | None) -> tuple[list[int], list[dict]]:
    """Recorre la primera fila del Excel, revisa el color de cada encabezado y devuelve cuáles columnas se deben conservar."""
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    columnas_conservar = []  # Guarda los índices de las columnas que sí deben quedarse
    resumen = []  # Guarda un registro simple de cada encabezado y si se conserva

    for cell in ws[1]:
        columna_excel = cell.column_letter
        indice_pandas = cell.column - 1
        nombre_columna = cell.value
        color = obtener_color_celda(cell)
        conservar = es_color_objetivo(color)

        if conservar:
            columnas_conservar.append(indice_pandas)  # Marca la columna para conservarla en la salida

        resumen.append({
            "columna_excel": columna_excel,
            "indice_pandas": indice_pandas,
            "nombre_columna": nombre_columna,
            "color_detectado": color,
            "conservar": conservar
        })

    wb.close()
    return columnas_conservar, resumen


def limpiar_excel(input_path: Path, sheet_name: str | None) -> None:
    """Abre el archivo, conserva solo las columnas marcadas y genera un archivo nuevo de salida."""
    if not input_path.exists():
        raise FileNotFoundError(
            f"No se encontró el archivo: {input_path}. "
            "Copia Data_proyecto_integrador.xlsx dentro de data/raw/."
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    logger.info("Inicio limpieza de columnas por color")
    logger.info("Archivo de entrada: {}", input_path)

    columnas_conservar, resumen = detectar_columnas_por_color(input_path, sheet_name)

    if not columnas_conservar:
        raise ValueError(
            "No se detectaron columnas con encabezado verde o anaranjado. "
            "Revisa que los colores estén aplicados en la primera fila del Excel."
        )

    df = pd.read_excel(input_path, sheet_name=sheet_name if sheet_name else 0)
    df_limpio = df.iloc[:, columnas_conservar].copy()  # Crea una copia solo con las columnas seleccionadas

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_excel = OUTPUT_DIR / f"{input_path.stem}_columnas_color_{timestamp}.xlsx"
    output_resumen = OUTPUT_DIR / "resumen_columnas_por_color.csv"

    df_limpio.to_excel(output_excel, index=False)  # Guarda el archivo limpio en Excel
    pd.DataFrame(resumen).to_csv(output_resumen, index=False, encoding="utf-8-sig")  # Guarda el resumen de columnas

    logger.info("Filas originales: {}", len(df))
    logger.info("Columnas originales: {}", len(df.columns))
    logger.info("Columnas conservadas: {}", len(df_limpio.columns))
    logger.info("Archivo limpio generado: {}", output_excel)
    logger.info("Resumen de columnas generado: {}", output_resumen)

    print("\nProceso finalizado correctamente.")
    print(f"Archivo limpio: {output_excel}")
    print(f"Resumen columnas: {output_resumen}")


def main() -> None:
    """Inicia el proceso desde la terminal y muestra mensajes si algo falla."""
    parser = argparse.ArgumentParser(
        description="Conserva columnas de Excel marcadas en verde o anaranjado."
    )
    parser.add_argument(
        "--input",
        default=str(INPUT_DEFAULT),
        help="Ruta del archivo Excel de entrada."
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Nombre de la hoja. Si se omite, usa la primera hoja."
    )
    args = parser.parse_args()

    log_path = configurar_logs()

    try:
        limpiar_excel(Path(args.input), args.sheet)
        logger.success("Proceso terminado correctamente")
        print(f"Log generado: {log_path}")
    except Exception as exc:
        logger.exception("Error durante la limpieza")
        print("\nOcurrió un error durante la ejecución.")
        print(f"Detalle: {exc}")
        print(f"Revisa el log: {log_path}")
        raise


if __name__ == "__main__":
    main()
