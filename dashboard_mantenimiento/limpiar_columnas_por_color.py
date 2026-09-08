"""
limpiar_columnas_por_color.py

Lee un archivo Excel y conserva solo las columnas cuyo encabezado esté marcado
con color verde o anaranjado. Genera salida en outputs/ y log en logs/.

Uso:
    python limpiar_columnas_por_color.py

Uso indicando archivo y hoja:
    python limpiar_columnas_por_color.py --input data/raw/Data_proyecto_integrador.xlsx --sheet Hoja1
"""

from pathlib import Path
from datetime import datetime
import argparse
import logging

import pandas as pd
from openpyxl import load_workbook


INPUT_DEFAULT = Path("data/raw/Data_proyecto_integrador.xlsx")
OUTPUT_DIR = Path("outputs")
LOG_DIR = Path("logs")


GREEN_HINTS = {
    "00B050", "92D050", "70AD47", "A9D18E", "C6E0B4",
    "008000", "00FF00"
}

ORANGE_HINTS = {
    "FFC000", "F4B183", "ED7D31", "FFA500", "FCE4D6",
    "FF9900", "FF6600"
}


def configurar_logging() -> Path:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = LOG_DIR / f"limpieza_columnas_color_{timestamp}.log"

    logging.basicConfig(
        filename=log_path,
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
        encoding="utf-8"
    )

    console = logging.StreamHandler()
    console.setLevel(logging.INFO)
    console.setFormatter(logging.Formatter("%(levelname)s | %(message)s"))
    logging.getLogger("").addHandler(console)

    return log_path


def normalizar_rgb(rgb: str | None) -> str:
    if rgb is None:
        return ""
    rgb = str(rgb).upper().replace("#", "")
    if len(rgb) == 8:
        rgb = rgb[-6:]
    return rgb


def es_color_objetivo(rgb: str) -> bool:
    rgb = normalizar_rgb(rgb)

    if not rgb:
        return False

    if rgb in GREEN_HINTS or rgb in ORANGE_HINTS:
        return True

    try:
        r = int(rgb[0:2], 16)
        g = int(rgb[2:4], 16)
        b = int(rgb[4:6], 16)
    except ValueError:
        return False

    es_verde = g >= 140 and g > r and g > b
    es_anaranjado = r >= 200 and 80 <= g <= 210 and b <= 120

    return es_verde or es_anaranjado


def obtener_color_celda(cell) -> str:
    fill = cell.fill

    if fill is None or fill.fill_type is None:
        return ""

    fg_color = fill.fgColor

    if fg_color.type == "rgb":
        return normalizar_rgb(fg_color.rgb)

    if fg_color.type == "indexed":
        return f"INDEXED_{fg_color.indexed}"

    if fg_color.type == "theme":
        return f"THEME_{fg_color.theme}"

    return ""


def detectar_columnas_por_color(excel_path: Path, sheet_name: str | None) -> tuple[list[int], list[dict]]:
    wb = load_workbook(excel_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]

    columnas_conservar = []
    resumen = []

    for cell in ws[1]:
        columna_excel = cell.column_letter
        indice_pandas = cell.column - 1
        nombre_columna = cell.value
        color = obtener_color_celda(cell)
        conservar = es_color_objetivo(color)

        if conservar:
            columnas_conservar.append(indice_pandas)

        resumen.append({
            "columna_excel": columna_excel,
            "indice_pandas": indice_pandas,
            "nombre_columna": nombre_columna,
            "color_detectado": color,
            "conservar": conservar
        })

    wb.close()
    return columnas_conservar, resumen


def limpiar_excel(input_path: Path, sheet_name: str | None) -> None:
    if not input_path.exists():
        raise FileNotFoundError(
            f"No se encontró el archivo: {input_path}. "
            "Copia Data_proyecto_integrador.xlsx dentro de data/raw/."
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    logging.info("Inicio limpieza de columnas por color")
    logging.info("Archivo de entrada: %s", input_path)

    columnas_conservar, resumen = detectar_columnas_por_color(input_path, sheet_name)

    if not columnas_conservar:
        raise ValueError(
            "No se detectaron columnas con encabezado verde o anaranjado. "
            "Revisa que los colores estén aplicados en la primera fila del Excel."
        )

    df = pd.read_excel(input_path, sheet_name=sheet_name if sheet_name else 0)
    df_limpio = df.iloc[:, columnas_conservar].copy()

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_excel = OUTPUT_DIR / f"{input_path.stem}_columnas_color_{timestamp}.xlsx"
    output_resumen = OUTPUT_DIR / "resumen_columnas_por_color.csv"

    df_limpio.to_excel(output_excel, index=False)
    pd.DataFrame(resumen).to_csv(output_resumen, index=False, encoding="utf-8-sig")

    logging.info("Filas originales: %s", len(df))
    logging.info("Columnas originales: %s", len(df.columns))
    logging.info("Columnas conservadas: %s", len(df_limpio.columns))
    logging.info("Archivo limpio generado: %s", output_excel)
    logging.info("Resumen de columnas generado: %s", output_resumen)

    print("\nProceso finalizado correctamente.")
    print(f"Archivo limpio: {output_excel}")
    print(f"Resumen columnas: {output_resumen}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Conserva columnas de Excel marcadas en verde o anaranjado."
    )
    parser.add_argument(
        "--input",
        default=str(INPUT_DEFAULT),
        help="Ruta del archivo Excel de entrada."
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Nombre de la hoja. Si se omite, usa la primera hoja."
    )
    args = parser.parse_args()

    log_path = configurar_logging()

    try:
        limpiar_excel(Path(args.input), args.sheet)
        logging.info("Proceso terminado correctamente")
        print(f"Log generado: {log_path}")
    except Exception as exc:
        logging.exception("Error durante la limpieza")
        print("\nOcurrió un error durante la ejecución.")
        print(f"Detalle: {exc}")
        print(f"Revisa el log: {log_path}")
        raise


if __name__ == "__main__":
    main()
